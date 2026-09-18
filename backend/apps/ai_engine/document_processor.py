import os
import io
import re
from typing import Tuple

class DocumentProcessor:
    """Extracts and normalizes text from multiple legal document formats."""

    @staticmethod
    def extract_text(file_path: str, file_type: str = '') -> Tuple[str, int, int]:
        """
        Extract text, page count, and word count from document.
        Returns: (extracted_text, page_count, word_count)
        """
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = file_type.lower() if file_type else file_path.split('.')[-1].lower()
        extracted_text = ""
        page_count = 1

        if ext == 'pdf':
            extracted_text, page_count = DocumentProcessor._extract_from_pdf(file_path)
        elif ext in ['docx', 'doc']:
            extracted_text, page_count = DocumentProcessor._extract_from_docx(file_path)
        elif ext in ['png', 'jpg', 'jpeg', 'tiff']:
            extracted_text, page_count = DocumentProcessor._extract_from_image(file_path)
        elif ext == 'txt':
            extracted_text, page_count = DocumentProcessor._extract_from_txt(file_path)
        else:
            # Fallback attempt plain text
            extracted_text, page_count = DocumentProcessor._extract_from_txt(file_path)

        cleaned_text = DocumentProcessor.clean_text(extracted_text)
        word_count = len(re.findall(r'\b\w+\b', cleaned_text))

        return cleaned_text, page_count, word_count

    @staticmethod
    def _extract_from_pdf(file_path: str) -> Tuple[str, int]:
        text_parts = []
        page_count = 0
        
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(file_path)
            page_count = len(doc)
            for page_num in range(page_count):
                page = doc.load_page(page_num)
                page_text = page.get_text()
                if page_text.strip():
                    text_parts.append(f"--- [Page {page_num + 1}] ---\n{page_text}")
                else:
                    # Try OCR on empty/scanned PDF page if pytesseract available
                    try:
                        import pytesseract
                        from PIL import Image
                        pix = page.get_pixmap(dpi=200)
                        img = Image.open(io.BytesIO(pix.tobytes("png")))
                        ocr_text = pytesseract.image_to_string(img)
                        if ocr_text.strip():
                            text_parts.append(f"--- [Page {page_num + 1} (OCR)] ---\n{ocr_text}")
                    except Exception:
                        pass
            doc.close()
        except ImportError:
            # Fallback to pypdf
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                page_count = len(reader.pages)
                for idx, page in enumerate(reader.pages):
                    pt = page.extract_text() or ''
                    text_parts.append(f"--- [Page {idx + 1}] ---\n{pt}")
            except Exception as e:
                text_parts.append(f"[Error extracting PDF: {str(e)}]")

        return "\n\n".join(text_parts), max(page_count, 1)

    @staticmethod
    def _extract_from_docx(file_path: str) -> Tuple[str, int]:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            text = "\n\n".join(paragraphs)
            # Estimate pages (approx 400 words per page)
            words = len(text.split())
            pages = max(1, (words // 400) + 1)
            return text, pages
        except Exception as e:
            return f"[Error extracting DOCX: {str(e)}]", 1

    @staticmethod
    def _extract_from_image(file_path: str) -> Tuple[str, int]:
        try:
            import pytesseract
            from PIL import Image
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img)
            return text, 1
        except Exception as e:
            return f"[OCR Image Extraction unavailable or failed: {str(e)}]", 1

    @staticmethod
    def _extract_from_txt(file_path: str) -> Tuple[str, int]:
        encodings = ['utf-8', 'latin-1', 'cp1252']
        for enc in encodings:
            try:
                with open(file_path, 'r', encoding=enc) as f:
                    text = f.read()
                    words = len(text.split())
                    pages = max(1, (words // 400) + 1)
                    return text, pages
            except (UnicodeDecodeError, Exception):
                continue
        return "", 1

    @staticmethod
    def clean_text(text: str) -> str:
        if not text:
            return ""
        # Remove extra whitespace while preserving paragraph structure
        text = re.sub(r'[ \t]+', ' ', text)
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()
